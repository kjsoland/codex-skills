#!/usr/bin/env python3
"""Extract common Sol Beam Delivery result tables and limiting values to JSON."""

from __future__ import annotations

import argparse
import json
import math
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def clean(value: Any) -> Any:
    if isinstance(value, float) and math.isfinite(value):
        return round(value, 9)
    return value


def number(value: Any) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return None


def find_row(ws, column: int, label: str) -> int | None:
    wanted = label.strip().casefold()
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row, column).value
        if isinstance(value, str) and value.strip().casefold() == wanted:
            return row
    return None


def find_row_prefix(ws, column: int, prefix: str) -> int | None:
    wanted = prefix.strip().casefold()
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row, column).value
        if isinstance(value, str) and value.strip().casefold().startswith(wanted):
            return row
    return None


def requirement_lookup(ws, collimator: str, label: str) -> Any:
    col = None
    for candidate in range(2, ws.max_column + 1):
        if ws.cell(1, candidate).value == collimator:
            col = candidate
            break
    row = find_row(ws, 1, label)
    return clean(ws.cell(row, col).value) if row and col else None


def extract_compliance(ws, requirements_ws) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    labels = {
        "wavelength_nm": "Test Wavelength [nm]",
        "position_x_um": "Relative Waist Lateral Position, X [um]",
        "position_y_um": "Relative Waist Lateral Position, Y [um]",
        "axial_u_mm": "Waist Axial Position Error u, Z [mm]",
        "axial_v_mm": "Waist Axial Position Error v, Z [mm]",
        "waist_u_um": "Waist Radius, u [um]",
        "waist_v_um": "Waist Radius, v [um]",
        "m2_u": "M2 Beam Quality, u",
        "m2_v": "M2 Beam Quality, v",
        "relative_pitch_mrad": "Relative Output Angle, Pitch",
        "relative_yaw_mrad": "Relative Output Angle, Yaw",
        "absolute_pitch_mrad": "Absolute Output Angle, Pitch",
        "absolute_yaw_mrad": "Absolute Output Angle, Yaw",
        "transmission": "Power Transmission",
    }
    rows: dict[str, int | None] = {}
    for key, label in labels.items():
        rows[key] = find_row(ws, 1, label) or find_row_prefix(ws, 1, label)

    records: list[dict[str, Any]] = []
    for col in range(2, min(ws.max_column, 9) + 1):
        name = ws.cell(1, col).value
        if not name:
            continue
        record: dict[str, Any] = {
            "collimator": name,
            "designator": ws.cell(2, col).value,
        }
        for key, row in rows.items():
            record[key] = clean(ws.cell(row, col).value) if row else None
        record["waist_target_um"] = requirement_lookup(
            requirements_ws, str(name), "Waist Radius [um]"
        )
        record["m2_limit"] = requirement_lookup(
            requirements_ws, str(name), "M2 Beam Quality (<)"
        )
        records.append(record)

    extrema: dict[str, Any] = {}
    axial = []
    waist = []
    m2 = []
    transmission = []
    for record in records:
        for axis in ("u", "v"):
            axial_value = number(record.get(f"axial_{axis}_mm"))
            if axial_value is not None:
                axial.append((abs(axial_value), axial_value, record["collimator"], axis))
            waist_value = number(record.get(f"waist_{axis}_um"))
            if waist_value is not None:
                target = number(record.get("waist_target_um"))
                deviation = abs(waist_value - target) if target is not None else None
                waist.append((waist_value, deviation, record["collimator"], axis, target))
            m2_value = number(record.get(f"m2_{axis}"))
            if m2_value is not None:
                m2.append((m2_value, record["collimator"], axis))
        transmission_value = number(record.get("transmission"))
        if transmission_value is not None:
            transmission.append((transmission_value, record["collimator"]))

    if axial:
        _, signed, coll, axis = max(axial)
        extrema["axial_worst"] = {
            "collimator": coll,
            "axis": axis,
            "signed_mm": clean(signed),
            "signed_um": clean(signed * 1000),
            "absolute_um": clean(abs(signed) * 1000),
        }
    if waist:
        low = min(waist, key=lambda item: item[0])
        high = max(waist, key=lambda item: item[0])
        with_target = [item for item in waist if item[1] is not None]
        extrema["waist_min"] = {
            "value_um": clean(low[0]), "collimator": low[2], "axis": low[3]
        }
        extrema["waist_max"] = {
            "value_um": clean(high[0]), "collimator": high[2], "axis": high[3]
        }
        if with_target:
            worst = max(with_target, key=lambda item: item[1])
            extrema["waist_worst_deviation"] = {
                "value_um": clean(worst[0]),
                "target_um": clean(worst[4]),
                "absolute_deviation_um": clean(worst[1]),
                "collimator": worst[2],
                "axis": worst[3],
            }
    if m2:
        worst = max(m2)
        extrema["m2_worst"] = {
            "value": clean(worst[0]), "collimator": worst[1], "axis": worst[2]
        }
    if transmission:
        worst = min(transmission)
        extrema["transmission_min"] = {
            "value": clean(worst[0]),
            "percent": clean(worst[0] * 100),
            "collimator": worst[1],
            "loss_db": clean(-10 * math.log10(worst[0])) if worst[0] > 0 else None,
        }
    return records, extrema


def extract_lateral(ws) -> dict[str, Any]:
    result: dict[str, Any] = {"individual": {}, "global": {}}
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row, 3).value
        if not isinstance(label, str):
            continue
        normalized = label.strip().casefold()
        if "range from nominal" not in normalized:
            continue
        values = [
            clean(ws.cell(row, col).value)
            for col in range(4, min(ws.max_column, 11) + 1)
            if ws.cell(row, col).value is not None
        ]
        key = (
            normalized.replace(" from nominal", "")
            .replace(" [um]", "")
            .replace(", ", "_")
            .replace(" ", "_")
        )
        group = "global" if ws.cell(row, 2).value == "Global" or row < 17 else "individual"
        result[group][key] = values
    return result


def extract_power(ws) -> list[dict[str, Any]]:
    rows = {
        "laser_on": find_row_prefix(ws, 1, "Laser On, at ion plane"),
        "laser_off": find_row_prefix(ws, 1, "Laser Off, at ion plane"),
        "out_of_fiber": find_row(ws, 1, "Out of Fiber"),
        "transmission": find_row_prefix(ws, 1, "Transmission"),
    }
    header_row = find_row(ws, 1, "Zone")
    if not header_row or not any(rows.values()):
        return []
    records = []
    for col in range(2, ws.max_column + 1):
        name = ws.cell(header_row, col).value
        if not name:
            continue
        record = {"collimator": name}
        for key, row in rows.items():
            record[key] = clean(ws.cell(row, col).value) if row else None
        if any(number(record[key]) is not None for key in rows):
            records.append(record)
    return records


def extract_polarization(ws) -> list[dict[str, Any]]:
    section = find_row(ws, 1, "Ion Plane Polarization")
    if not section:
        return []
    records = []
    for row in range(section + 1, ws.max_row + 1):
        name = ws.cell(row, 1).value
        if not isinstance(name, str) or not name.startswith("Col_"):
            if records and row > section + 20:
                break
            continue
        per = ws.cell(row, 2).value
        state = ws.cell(row, 3).value
        ellipticity = ws.cell(row, 4).value
        azimuth = ws.cell(row, 5).value
        if any(value is not None for value in (per, state, ellipticity, azimuth)):
            records.append(
                {
                    "collimator": name,
                    "per": clean(per),
                    "state": state,
                    "ellipticity": clean(ellipticity),
                    "azimuth": clean(azimuth),
                }
            )
    return records


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if not args.workbook.is_file():
        parser.error(f"Workbook not found: {args.workbook}")

    wb = load_workbook(args.workbook, read_only=False, data_only=True)
    required = {"Compliance", "Requirements", "Lateral Range", "Data Input"}
    missing = sorted(required - set(wb.sheetnames))
    if missing:
        raise SystemExit("Missing expected sheets: " + ", ".join(missing))

    compliance, extrema = extract_compliance(wb["Compliance"], wb["Requirements"])
    port = wb["Requirements"]["A36"].value
    result = {
        "source": str(args.workbook.resolve()),
        "port": port,
        "compliance_records": compliance,
        "extrema": extrema,
        "lateral_range": extract_lateral(wb["Lateral Range"]),
        "power_records": extract_power(wb["Data Input"]),
        "polarization_records": extract_polarization(wb["Data Input"]),
    }
    wb.close()

    rendered = json.dumps(result, indent=2, ensure_ascii=False)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        if args.output.exists():
            raise SystemExit(f"Refusing to overwrite output: {args.output}")
        args.output.write_text(rendered + "\n", encoding="utf-8")
        print(args.output)
    else:
        sys.stdout.write(rendered + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
