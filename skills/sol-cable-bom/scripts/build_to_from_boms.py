#!/usr/bin/env python3
"""Build the remaining Sol To/From ICD cable-kit BOM workbooks."""

from __future__ import annotations

import argparse
import json
import re
from collections import OrderedDict
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font


HEADERS = [
    "Item No.",
    "Arena Part Number",
    "Vendor Part Number",
    "Quantity",
    "Description",
    "Notes",
]


@dataclass(frozen=True)
class BomLine:
    vendor_part_number: str
    quantity: int | float
    description: str


@dataclass(frozen=True)
class Profile:
    name: str
    icd_folder: str
    source_pattern: str
    output_folder: str
    output_name: str
    extractor: Callable[[Path], tuple[list[BomLine], int]]


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def clean_length(value: Any) -> str:
    value_text = text(value).lstrip("=").strip()
    if not value_text or value_text in {"-", "NONE"}:
        return ""
    try:
        number = float(value_text)
        return str(int(number)) if number.is_integer() else f"{number:g}"
    except ValueError:
        return value_text


def rows(path: Path, sheet_name: str, *, data_only: bool = True) -> list[tuple[int, dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=data_only)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"{path.name}: missing worksheet {sheet_name!r}")
    worksheet = workbook[sheet_name]
    values = worksheet.iter_rows(values_only=True)
    headers = [text(value) for value in next(values)]
    result = [(row_number, dict(zip(headers, values_row))) for row_number, values_row in enumerate(values, start=2)]
    workbook.close()
    return result


def grouped_lines(
    records: Iterable[dict[str, Any]],
    key: Callable[[dict[str, Any]], tuple[Any, ...]],
    line: Callable[[dict[str, Any], int], BomLine],
) -> list[BomLine]:
    groups: OrderedDict[tuple[Any, ...], list[dict[str, Any]]] = OrderedDict()
    for record in records:
        groups.setdefault(key(record), []).append(record)
    return [line(group[0], len(group)) for group in groups.values()]


def length_phrase(length: str) -> str:
    return f", {length} in" if length else "; length TBD"


def extract_ac_power(path: Path) -> tuple[list[BomLine], int]:
    source: list[dict[str, Any]] = []
    for _row, record in rows(path, "Electrical Cabling"):
        cable_type = text(record.get("Cable Type"))
        if cable_type not in {"Power", "Ground"}:
            continue
        if not any(text(record.get(name)) for name in ("From Component", "To Component")):
            continue
        source.append(
            {
                "type": cable_type,
                "length": clean_length(record.get("Length (inches)")),
                "source_part": text(record.get("Cable Part Number")),
            }
        )

    def make(record: dict[str, Any], quantity: int) -> BomLine:
        cable_type = record["type"]
        length = record["length"]
        source_part = record["source_part"]
        detail = ""
        if source_part:
            detail = f"; source calls for {source_part}"
        description = f"{cable_type} cable assembly{length_phrase(length)}{detail}; vendor part selection required"
        return BomLine("TBD", quantity, description)

    lines = grouped_lines(source, lambda item: (item["type"], item["length"], item["source_part"]), make)
    return lines, len(source)


def extract_bd_beamplate(path: Path) -> tuple[list[BomLine], int]:
    source = []
    for row_number, record in rows(path, "Electrical Cabling"):
        if row_number > 174:
            break
        part = text(record.get("Cable Part Number"))
        if not part or part.startswith("="):
            continue
        source.append(
            {
                "part": part,
                "vendor": text(record.get("Vendor")),
                "length": clean_length(record.get("Length (Inches)")),
            }
        )

    def make(record: dict[str, Any], quantity: int) -> BomLine:
        part = record["part"]
        length = record["length"]
        if part.startswith("410258"):
            description = f"Micronix 26-pin Gecko female-to-female extension cable{length_phrase(length)}"
        else:
            meters = "10 m" if part.endswith("01-EPLSP") else "12 m"
            description = f"Samtec EPLSP-019 rugged high-speed cable assembly, {meters} custom length ({length} in)"
        return BomLine(part, quantity, description)

    lines = grouped_lines(source, lambda item: (item["part"], item["length"]), make)
    return lines, len(source)


def extract_cooling(path: Path) -> tuple[list[BomLine], int]:
    source = []
    for _row, record in rows(path, "Water Tubes"):
        part = text(record.get("Cable Part Number"))
        description = text(record.get("Description"))
        length = clean_length(record.get("Length (inches)"))
        if not part or part.startswith("=") or not length:
            continue
        if not any(text(record.get(name)) for name in ("From Component", "To Component")):
            continue
        source.append(
            {
                "part": part,
                "vendor": text(record.get("Vendor")),
                "description": description,
                "length": Decimal(length),
            }
        )

    groups: OrderedDict[tuple[str, str], list[dict[str, Any]]] = OrderedDict()
    for record in source:
        groups.setdefault((record["part"], record["description"].casefold()), []).append(record)

    lines: list[BomLine] = []
    for group in groups.values():
        total_feet = sum((record["length"] for record in group), Decimal("0")) / Decimal("12")
        total_feet = total_feet.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        quantity = int(total_feet) if total_feet == total_feet.to_integral_value() else float(total_feet)
        first = group[0]
        description = f"{first['vendor']} {first['description']}; quantity is in feet"
        incompatible_descriptions = {
            record["description"].casefold()
            for record in source
            if record["part"] == first["part"]
        }
        if len(incompatible_descriptions) > 1:
            description += f"; source assigns {first['part']} to multiple colors—confirm the color-specific vendor part"
            lines.append(BomLine("TBD", quantity, description))
        else:
            lines.append(BomLine(first["part"], quantity, description))
    return lines, len(source)


def extract_core_control(path: Path) -> tuple[list[BomLine], int]:
    source: list[dict[str, Any]] = []
    for sheet_name in ("DDS Channel Map", "ADC Channel Map"):
        for _row, record in rows(path, sheet_name):
            if text(record.get("External Cable Present?")).casefold() != "yes":
                continue
            part = text(record.get("Cable Part Number"))
            if not part or part == "-":
                continue
            source.append(
                {
                    "part": part,
                    "vendor": text(record.get("Cable Vendor")),
                    "length": clean_length(record.get("Cable Length (in)")),
                    "type": "RF coaxial cable assembly",
                }
            )

    for _row, record in rows(path, "TTL Channel Map"):
        part = text(record.get("Cable Part Number"))
        if not part or part == "-":
            continue
        source.append(
            {
                "part": part,
                "vendor": text(record.get("Cable Vendor")),
                "length": clean_length(record.get("Cable Length (in)")),
                "type": "TTL cable assembly",
            }
        )

    for _row, record in rows(path, "PMT Channel Map"):
        part_cell = text(record.get("Cable Part Number"))
        if not part_cell:
            continue
        part_tokens = [token.strip() for token in re.split(r"[\r\n]+", part_cell) if token.strip()]
        for part in part_tokens:
            source.append(
                {
                    "part": part,
                    "vendor": text(record.get("Cable Vendor")),
                    "length": "6 m",
                    "type": "custom VHDCI PMT cable assembly",
                }
            )

    def make(record: dict[str, Any], quantity: int) -> BomLine:
        length = record["length"]
        length_text = f", {length}" if length.endswith("m") else length_phrase(length)
        description = f"{record['vendor']} {record['type']}{length_text}"
        return BomLine(record["part"], quantity, description)

    lines = grouped_lines(source, lambda item: (item["part"], item["length"], item["type"]), make)
    return lines, len(source)


def extract_electrode(path: Path) -> tuple[list[BomLine], int]:
    source = []
    for _row, record in rows(path, "Cabling"):
        part = text(record.get("Cable PN"))
        if not part or part.casefold() == "none":
            continue
        length = clean_length(record.get("Length"))
        source.append(
            {
                "part": "TBD" if part.casefold() == "tbd" else part,
                "vendor": text(record.get("Vendor")),
                "length": length,
            }
        )

    def make(record: dict[str, Any], quantity: int) -> BomLine:
        if record["part"] == "TBD":
            description = "QTM electrode cable assembly; length and vendor part number TBD"
        else:
            description = f"Samtec C28S electrode cable assembly, {record['length']}"
        return BomLine(record["part"], quantity, description)

    lines = grouped_lines(source, lambda item: (item["part"], item["length"]), make)
    return lines, len(source)


def extract_frequency(path: Path) -> tuple[list[BomLine], int]:
    source = []
    for row_number, record in rows(path, "Electrical Cabling_All"):
        if row_number > 119 or text(record.get("External / Internal")).casefold() != "external":
            continue
        part = text(record.get("Cable Part Number"))
        if not part or part.startswith("See "):
            continue
        source.append(
            {
                "part": part,
                "vendor": text(record.get("Vendor")),
                "length": clean_length(record.get("Length")),
                "description": text(record.get("Cable Description")),
            }
        )

    def make(record: dict[str, Any], quantity: int) -> BomLine:
        description = record["description"] or "External frequency-reference cable assembly"
        if record["length"]:
            description = f"{description}, {record['length']} in"
        else:
            description = f"{description}; length TBD"
        part = record["part"]
        if "#" in part:
            description += f"; source part family {part} requires an exact length/part selection"
            part = "TBD"
        return BomLine(part, quantity, description)

    lines = grouped_lines(source, lambda item: (item["part"], item["length"], item["description"]), make)
    return lines, len(source)


def extract_gates_lsc(path: Path) -> tuple[list[BomLine], int]:
    source = []
    for row_number, record in rows(path, "Electrical Cabling"):
        if row_number > 119:
            break
        part = text(record.get("Cable Part Number"))
        if not part or part.startswith("="):
            continue
        source.append(
            {
                "part": part,
                "vendor": text(record.get("Vendor")),
                "length": clean_length(record.get("Length")),
            }
        )

    def make(record: dict[str, Any], quantity: int) -> BomLine:
        description = f"{record['vendor']} RF cable assembly{length_phrase(record['length'])}"
        part = record["part"]
        if record["part"].endswith("-315") and record["length"] == "354":
            description += f"; ICD length conflicts with source part {record['part']}—confirm the correct vendor part before release"
            part = "TBD"
        return BomLine(part, quantity, description)

    lines = grouped_lines(source, lambda item: (item["part"], item["length"]), make)
    return lines, len(source)


def extract_qpa(path: Path) -> tuple[list[BomLine], int]:
    candidates: list[dict[str, Any]] = []
    selected_connectors: set[tuple[str, str, str]] = set()
    for row_number, record in rows(path, "Electrical Cabling"):
        if row_number > 73:
            break
        from_location = text(record.get("From Location"))
        to_location = text(record.get("To Location"))
        if not from_location or not to_location or from_location.startswith("See "):
            continue
        length = clean_length(record.get("Length \n(Inches)"))
        vendor = text(record.get("Vendor"))
        raw_part = text(record.get("Cable Part Number"))
        source_description = text(record.get("Cable Description"))
        if all(value == "-" for value in (text(record.get("Length \n(Inches)")), vendor, raw_part, source_description)):
            continue
        connector = (from_location, text(record.get("From Component")), text(record.get("From Connection")))
        has_purchase_cue = bool(length or vendor or raw_part or source_description)
        if not has_purchase_cue and connector in selected_connectors:
            continue
        selected_connectors.add(connector)
        designator = text(record.get("Signal Designator"))
        part = raw_part
        if not part or part.casefold() in {"tbd", "need a coupler"}:
            part = "TBD"

        if source_description and source_description != "-":
            description = source_description.rstrip()
        elif designator.casefold().startswith("mech power"):
            description = "QTM MECH power cable assembly"
        elif designator == "AntennaDriver_DCPWR":
            description = "Antenna-driver DC power cable assembly"
        elif designator == "SPAD_Power":
            description = "SPAD power cable assembly"
        elif raw_part.casefold() == "need a coupler":
            description = "Motor cable coupler; vendor part selection required"
        elif designator.startswith("TS3_"):
            description = "QPA temperature/heater cable assembly"
        else:
            description = f"{vendor + ' ' if vendor else ''}cable assembly".strip()

        if length:
            description += f", {length} in"
        else:
            description += "; length TBD"
        if part == "TBD" and "vendor part" not in description.casefold():
            description += "; vendor part selection required"
        candidates.append({"part": part, "length": length, "description": description})

    lines = grouped_lines(
        candidates,
        lambda item: (item["part"], item["length"], item["description"]),
        lambda item, quantity: BomLine(item["part"], quantity, item["description"]),
    )
    return lines, len(candidates)


def extract_fiber(path: Path) -> tuple[list[BomLine], int]:
    source = []
    for _row, record in rows(path, "External Fibers", data_only=True):
        if not text(record.get("From Component")) or not text(record.get("To Component")):
            continue
        part = text(record.get("QTM Part Number"))
        if not part:
            continue
        source.append(
            {
                "part": part,
                "type": text(record.get("Fiber Type")),
                "connector_a": text(record.get("Connector A")),
                "connector_b": text(record.get("Connector B")),
                "length": clean_length(record.get("Fiber Length (m)")),
            }
        )

    def make(record: dict[str, Any], quantity: int) -> BomLine:
        description = (
            f"QTM {record['type']} fiber assembly, {record['connector_a']} to {record['connector_b']}, "
            f"{record['length']} m"
        )
        return BomLine(record["part"], quantity, description)

    lines = grouped_lines(source, lambda item: (item["part"],), make)
    return lines, len(source)


def extract_system_monitor(path: Path) -> tuple[list[BomLine], int]:
    source: list[dict[str, Any]] = []
    for row_number, record in rows(path, "Electrical Cabling"):
        if row_number > 115:
            break
        if not text(record.get("From Component")) or not text(record.get("To Component")):
            continue
        designator = text(record.get("Designator"))
        part = text(record.get("Cable Part Number"))
        vendor = text(record.get("Vendor"))
        length = clean_length(record.get("Length"))

        if designator.startswith("MECH_") and not part:
            continue
        if part:
            category = "system-monitor cable assembly"
            normalized_part = part
        elif vendor == "NI":
            category = "NI photodiode-breakout cable assembly"
            normalized_part = "TBD"
        elif designator.startswith("TS_"):
            category = "temperature-sensor cable assembly"
            normalized_part = "TBD"
        elif designator.startswith("Vibe_"):
            category = "accelerometer cable assembly"
            normalized_part = "TBD"
        elif designator.startswith(("DPLX", "PD_")):
            category = "photodiode signal cable assembly"
            normalized_part = "TBD"
        else:
            category = "system-monitor cable assembly"
            normalized_part = "TBD"

        source.append(
            {
                "part": normalized_part,
                "vendor": vendor,
                "length": length,
                "category": category,
            }
        )

    def make(record: dict[str, Any], quantity: int) -> BomLine:
        vendor_prefix = f"{record['vendor']} " if record["vendor"] and not record["category"].startswith(record["vendor"]) else ""
        description = f"{vendor_prefix}{record['category']}{length_phrase(record['length'])}"
        if record["part"] == "TBD":
            description += "; vendor part selection required"
        return BomLine(record["part"], quantity, description)

    lines = grouped_lines(source, lambda item: (item["part"], item["length"], item["category"]), make)
    return lines, len(source)


PROFILES = [
    Profile("AC Power and Ground", "ACPower-Grounding ICD", "AC_Power_Ground_ICD_v*.xlsx", "AC Power and Ground Cable Kit", "AC Power and Ground Cable Kit BOM v0p1.xlsx", extract_ac_power),
    Profile("BD Beamplate", "BD Beamplate Electrical ICD", "BD_Beamplate_EICD_v*.xlsx", "BD Beamplate Cable Kit", "BD Beamplate Cable Kit BOM v0p1.xlsx", extract_bd_beamplate),
    Profile("Cooling", "Cooling ICD", "Sol_Cooling_ICD_v*.xlsx", "Cooling Interconnect Kit", "Cooling Interconnect Kit BOM v0p1.xlsx", extract_cooling),
    Profile("Core Control", "Core Control Signal Map ICD", "Core Control Signal Map ICD v*.xlsx", "Core Control System Cable Kit", "Core Control System Cable Kit BOM v0p1.xlsx", extract_core_control),
    Profile("Electrode", "Electrode Signal Map ICD", "Electrode Signal Map ICD*.xlsx", "Electrode Cable Kit", "Electrode Cable Kit BOM v0p1.xlsx", extract_electrode),
    Profile("Fiber Optic", "Sol Fiber ICD", "Sol_Fiber_ICD_v*.xlsx", "Fiber Optic Kit", "Fiber Optic Kit BOM v0p1.xlsx", extract_fiber),
    Profile("Frequency Reference", "Frequency Reference ICD", "Frequency_Reference_ICD_v*.xlsx", "Frequency Reference Cable Kit", "Frequency Reference Cable Kit BOM v0p1.xlsx", extract_frequency),
    Profile("Gates, LSC Electrical", "Gates-LSC Electrical ICD", "Gates_LSC_Electrical_ICD_v*.xlsx", "Gates, LSC Electrical Cable Kit", "Gates, LSC Electrical Cable Kit BOM v0p1.xlsx", extract_gates_lsc),
    Profile("QPA and Physics Package Control", "QPA And Physics Package Control ICD", "QPA_and_Physics_Package_Control_ICD_v*.xlsx", "QPA and Physics Package Control Cable Kit", "QPA and Physics Package Control Cable Kit BOM v0p1.xlsx", extract_qpa),
    Profile("System Monitor", "System Monitor ICD", "Sol_System_Monitor_ICD_v*.xlsx", "System Monitor Cable Kit", "System Monitor Cable Kit BOM v0p1.xlsx", extract_system_monitor),
]


def version_key(path: Path) -> tuple[int, int, int, float]:
    version_matches = re.findall(r"(?i)v(\d+)(?:p(\d+))?", path.stem)
    major, minor = (int(version_matches[-1][0]), int(version_matches[-1][1] or 0)) if version_matches else (0, 0)
    recovery = 1 if "recovery" in path.stem.casefold() else 0
    return major, minor, recovery, path.stat().st_mtime


def resolve_source(icd_root: Path, profile: Profile, override: Path | None = None) -> Path:
    if override is not None:
        if not override.is_file():
            raise FileNotFoundError(override)
        return override
    folder = icd_root / profile.icd_folder
    direct = [path for path in folder.glob(profile.source_pattern) if not path.name.startswith("~$")]
    candidates = direct or [path for path in folder.rglob(profile.source_pattern) if not path.name.startswith("~$")]
    if not candidates:
        raise FileNotFoundError(f"No source matching {profile.source_pattern!r} under {folder}")
    return max(candidates, key=version_key)


def write_workbook(output: Path, title: str, source: Path, lines: list[BomLine]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Cable BOM"
    workbook.properties.title = f"Sol {title} Cable BOM"
    workbook.properties.subject = f"Cable BOM derived from {source.name}"
    workbook.properties.description = "Arena part numbers and drawing-note cells are intentionally blank."

    worksheet.append(HEADERS)
    for item_number, line in enumerate(lines, start=1):
        worksheet.append([item_number, None, line.vendor_part_number, line.quantity, line.description, None])

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 24

    for row_number in range(2, worksheet.max_row + 1):
        for column in range(1, 7):
            worksheet.cell(row_number, column).alignment = Alignment(vertical="top", wrap_text=column == 5)
        worksheet.cell(row_number, 1).alignment = Alignment(horizontal="center", vertical="top")
        worksheet.cell(row_number, 4).alignment = Alignment(horizontal="center", vertical="top")
        description_length = len(text(worksheet.cell(row_number, 5).value))
        worksheet.row_dimensions[row_number].height = min(45, max(20, 15 * (1 + description_length // 60)))

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:F{worksheet.max_row}"
    for column, width in {"A": 13, "B": 22, "C": 28, "D": 13, "E": 72, "F": 30}.items():
        worksheet.column_dimensions[column].width = width
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_title_rows = "1:1"
    worksheet.print_area = f"A1:F{worksheet.max_row}"

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def validate_workbook(output: Path, expected_lines: int) -> dict[str, Any]:
    workbook = load_workbook(output, data_only=False)
    if workbook.sheetnames != ["Cable BOM"]:
        raise AssertionError(f"{output.name}: unexpected worksheets {workbook.sheetnames}")
    worksheet = workbook["Cable BOM"]
    actual_headers = [worksheet.cell(1, column).value for column in range(1, 7)]
    if actual_headers != HEADERS:
        raise AssertionError(f"{output.name}: unexpected headers {actual_headers}")
    if worksheet.max_row - 1 != expected_lines:
        raise AssertionError(f"{output.name}: expected {expected_lines} lines, found {worksheet.max_row - 1}")
    if worksheet.freeze_panes != "A2" or worksheet.auto_filter.ref != f"A1:F{worksheet.max_row}":
        raise AssertionError(f"{output.name}: filter or freeze-pane validation failed")
    if worksheet.tables:
        raise AssertionError(f"{output.name}: workbook must not contain formatted tables")
    if not all(cell.font.bold for cell in worksheet[1]):
        raise AssertionError(f"{output.name}: header row must be bold")
    if any(worksheet.cell(row, 2).value not in (None, "") for row in range(2, worksheet.max_row + 1)):
        raise AssertionError(f"{output.name}: Arena Part Number must be blank")
    if any(worksheet.cell(row, 6).value not in (None, "") for row in range(2, worksheet.max_row + 1)):
        raise AssertionError(f"{output.name}: Notes must be blank")
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.fill.fill_type is not None or cell.hyperlink is not None or cell.comment is not None:
                raise AssertionError(f"{output.name}: found prohibited formatting or metadata")
            if isinstance(cell.value, str) and cell.value.startswith("="):
                raise AssertionError(f"{output.name}: output must not contain formulas")
    summary = {
        "output": str(output),
        "bom_lines": worksheet.max_row - 1,
        "total_quantity": sum(float(worksheet.cell(row, 4).value) for row in range(2, worksheet.max_row + 1)),
        "tbd_lines": sum(worksheet.cell(row, 3).value == "TBD" for row in range(2, worksheet.max_row + 1)),
    }
    workbook.close()
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("icd_root", type=Path, help="SharePoint-backed Systems/ICDs folder")
    parser.add_argument("product_root", type=Path, help="SharePoint-backed BOMs/Sol Computer/System Interconnect folder")
    parser.add_argument("--system-monitor-xlsx", type=Path, help="Readable .xlsx copy of the latest System Monitor .xlsb")
    args = parser.parse_args()

    summaries = []
    for profile in PROFILES:
        override = args.system_monitor_xlsx if profile.name == "System Monitor" else None
        source = resolve_source(args.icd_root, profile, override)
        lines, source_instances = profile.extractor(source)
        if not lines:
            raise AssertionError(f"{profile.name}: extraction produced no BOM lines")
        output = args.product_root / profile.output_folder / profile.output_name
        write_workbook(output, profile.name, source, lines)
        summary = validate_workbook(output, len(lines))
        summary.update({"profile": profile.name, "source": str(source), "source_instances": source_instances})
        summaries.append(summary)
    print(json.dumps(summaries, indent=2))


if __name__ == "__main__":
    main()
