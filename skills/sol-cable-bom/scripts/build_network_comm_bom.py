#!/usr/bin/env python3
"""Build the Sol Network and Comm cable-kit BOM from its ICD workbook."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font


HEADERS = [
    "Item No.",
    "Arena Part Number",
    "Vendor Part Number",
    "Quantity",
    "Description",
    "Notes",
]


@dataclass(frozen=True)
class CableRecord:
    sheet: str
    row: int
    designator: str
    cable_type: str
    length: str
    vendor: str
    part: str
    order_url: str
    source_note: str
    endpoints: str


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalized_length(value: Any) -> str:
    raw = text(value)
    return re.sub(r"(?i)^(\d+(?:\.\d+)?)(m|ft)$", r"\1 \2", raw)


def header_map(ws) -> dict[str, int]:
    return {text(cell.value): cell.column for cell in ws[1] if text(cell.value)}


def require_headers(ws, required: Iterable[str]) -> dict[str, int]:
    headers = header_map(ws)
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError(f"{ws.title}: missing required columns: {', '.join(missing)}")
    return headers


def cell(ws, row: int, headers: dict[str, int], name: str) -> Any:
    return ws.cell(row, headers[name]).value


def endpoint(*parts: Any) -> str:
    values = [text(value) for value in parts if text(value)]
    return " / ".join(values) if values else "TBD"


def compact_rows(rows: Iterable[int]) -> str:
    values = sorted(set(rows))
    if not values:
        return ""
    chunks: list[str] = []
    start = previous = values[0]
    for value in values[1:]:
        if value == previous + 1:
            previous = value
            continue
        chunks.append(str(start) if start == previous else f"{start}-{previous}")
        start = previous = value
    chunks.append(str(start) if start == previous else f"{start}-{previous}")
    return ", ".join(chunks)


def load_catalog() -> dict[str, Any]:
    path = Path(__file__).resolve().parents[1] / "references" / "network-comm-catalog.json"
    return json.loads(path.read_text(encoding="utf-8"))


def read_interrack(ws) -> list[CableRecord]:
    required = [
        "Designator",
        "Notes",
        "From Location",
        "From Component",
        "From Connection",
        "To Location",
        "To Component",
        "To Connection",
        "Cable Type",
        "Actual length (m)",
        "Vendor",
        "Cable Part Number",
    ]
    headers = require_headers(ws, required)
    records: list[CableRecord] = []
    for row in range(2, ws.max_row + 1):
        designator = text(cell(ws, row, headers, "Designator"))
        if not designator:
            continue
        part = text(cell(ws, row, headers, "Cable Part Number"))
        order_url = part if part.lower().startswith(("http://", "https://")) else ""
        records.append(
            CableRecord(
                sheet=ws.title,
                row=row,
                designator=designator,
                cable_type=text(cell(ws, row, headers, "Cable Type")) or "TBD",
                length=normalized_length(cell(ws, row, headers, "Actual length (m)")) or "TBD",
                vendor=text(cell(ws, row, headers, "Vendor")) or "TBD",
                part=part,
                order_url=order_url,
                source_note=text(cell(ws, row, headers, "Notes")),
                endpoints=(
                    f"{endpoint(cell(ws, row, headers, 'From Location'), cell(ws, row, headers, 'From Component'), cell(ws, row, headers, 'From Connection'))}"
                    f" -> {endpoint(cell(ws, row, headers, 'To Location'), cell(ws, row, headers, 'To Component'), cell(ws, row, headers, 'To Connection'))}"
                ),
            )
        )
    return records


def read_non_ethernet(ws) -> list[CableRecord]:
    required = [
        "Description",
        "Type",
        "Plan",
        "From Location",
        "From Component",
        "From Connection",
        "To Location",
        "To Component",
        "To Connection",
        "External Cable Present?",
        "Cable Length (in)",
        "Cable Vendor",
        "Cable Part Number",
    ]
    headers = require_headers(ws, required)
    records: list[CableRecord] = []
    for row in range(2, ws.max_row + 1):
        if text(cell(ws, row, headers, "External Cable Present?")).lower() != "yes":
            continue
        length_in = text(cell(ws, row, headers, "Cable Length (in)"))
        plan = text(cell(ws, row, headers, "Plan")) or "TBD"
        records.append(
            CableRecord(
                sheet=ws.title,
                row=row,
                designator=text(cell(ws, row, headers, "Description")) or "TBD",
                cable_type=text(cell(ws, row, headers, "Type")) or "TBD",
                length=f"{length_in} in" if length_in else "TBD",
                vendor=text(cell(ws, row, headers, "Cable Vendor")) or "TBD",
                part=text(cell(ws, row, headers, "Cable Part Number")),
                order_url="",
                source_note=f"Plan: {plan}",
                endpoints=(
                    f"{endpoint(cell(ws, row, headers, 'From Location'), cell(ws, row, headers, 'From Component'), cell(ws, row, headers, 'From Connection'))}"
                    f" -> {endpoint(cell(ws, row, headers, 'To Location'), cell(ws, row, headers, 'To Component'), cell(ws, row, headers, 'To Connection'))}"
                ),
            )
        )
    return records


def unique(values: Iterable[str]) -> list[str]:
    return sorted({value for value in values if value and value != "TBD"}, key=str.casefold)


def source_note(source_name: str, records: list[CableRecord]) -> str:
    rows = compact_rows(record.row for record in records)
    return f"Source: {source_name}, {records[0].sheet} row(s) {rows}."


def build_lines(source_name: str, interrack: list[CableRecord], non_ethernet: list[CableRecord], catalog: dict[str, Any]) -> list[dict[str, Any]]:
    link_catalog = catalog["links"]
    pn_catalog = catalog["part_numbers"]
    complete: dict[tuple[str, str, str], list[CableRecord]] = defaultdict(list)
    missing: dict[tuple[str, str, str], list[CableRecord]] = defaultdict(list)

    for record in interrack:
        entry = link_catalog.get(record.order_url)
        identity = text(entry.get("vendor_part_number")) if entry else (record.part if record.part and not record.order_url else "")
        if identity:
            complete[(identity, record.cable_type, record.length)].append(record)
        else:
            missing[(record.cable_type, record.length, record.designator)].append(record)

    identity_types: dict[str, Counter[str]] = defaultdict(Counter)
    for (identity, cable_type, _length), records in complete.items():
        identity_types[identity][cable_type] += len(records)

    lines: list[dict[str, Any]] = []
    for (identity, cable_type, length), records in complete.items():
        entries = [link_catalog[record.order_url] for record in records if record.order_url in link_catalog]
        entries_by_part = [entry for entry in entries if text(entry.get("vendor_part_number")) == identity]
        entry = entries_by_part[0] if entries_by_part else (entries[0] if entries else {})
        dominant_type = identity_types[identity].most_common(1)[0][0]
        conflicting_type = len(identity_types[identity]) > 1 and cable_type != dominant_type
        source_vendors = unique(record.vendor for record in records)
        catalog_vendor = text(entry.get("vendor"))
        urls = unique(record.order_url for record in records)

        if conflicting_type:
            vendor_pn = "TBD"
            description = f"{length} {cable_type.lower()} patch cable - correct part selection required"
            notes = (
                f"Source conflict: these rows are {cable_type}, but their purchasing link resolves to "
                f"{catalog_vendor or 'a'} {dominant_type.lower()} item {identity}. Vendor part number intentionally left TBD pending ICD correction. "
            )
            hyperlink = ""
        else:
            vendor_pn = identity
            description = text(entry.get("description")) or f"{length} {cable_type.lower()} cable"
            identifier_type = text(entry.get("identifier_type"))
            notes = f"Manufacturer/vendor: {catalog_vendor or ', '.join(source_vendors) or 'TBD'}. "
            if identifier_type:
                notes += f"Vendor Part Number is the {identifier_type} resolved from the ICD purchasing link; confirm the manufacturer part number before Arena release. "
            hyperlink = text(entry.get("canonical_url")) or (urls[0] if urls else "")

        if catalog_vendor and source_vendors and any(v.casefold() != catalog_vendor.casefold() for v in source_vendors):
            notes += f"ICD Vendor value(s): {', '.join(source_vendors)}; resolved link identifies {catalog_vendor}. "
        if len(urls) > 1:
            notes += f"Grouped {len(urls)} ICD purchasing links that resolve to the same order identifier. "
        notes += source_note(source_name, records)
        if urls:
            notes += f" ICD order link(s): {'; '.join(urls)}."
        design_notes = unique(record.source_note for record in records)
        if design_notes:
            notes += f" ICD note(s): {'; '.join(design_notes)}."
        lines.append(
            {
                "order": (0, min(record.row for record in records)),
                "vendor_part_number": vendor_pn,
                "quantity": len(records),
                "description": description,
                "notes": notes,
                "hyperlink": hyperlink,
            }
        )

    for (_cable_type, _length, _designator), records in missing.items():
        record = records[0]
        missing_fields = []
        if record.length == "TBD":
            missing_fields.append("actual length")
        if record.vendor == "TBD":
            missing_fields.append("vendor")
        if not record.part:
            missing_fields.append("vendor part number")
        notes = (
            f"Manufacturer/vendor: TBD. Missing ICD field(s): {', '.join(missing_fields) or 'vendor part number'}. "
            f"Connections: {'; '.join(unique(r.endpoints for r in records))}. {source_note(source_name, records)}"
        )
        lines.append(
            {
                "order": (0, min(r.row for r in records)),
                "vendor_part_number": "TBD",
                "quantity": len(records),
                "description": f"{record.length if record.length != 'TBD' else 'Length TBD'} {record.cable_type.lower()} cable for {record.designator}",
                "notes": notes,
                "hyperlink": "",
            }
        )

    external_groups: dict[tuple[str, str, str, str], list[CableRecord]] = defaultdict(list)
    for record in non_ethernet:
        external_groups[(record.part or "TBD", record.cable_type, record.length, record.designator)].append(record)

    for (part, cable_type, length, designator), records in external_groups.items():
        entry = pn_catalog.get(part, {})
        vendor_pn = part if part != "TBD" else "TBD"
        if part == "TBD":
            description = f"{cable_type} cable for {designator} - correct part selection required"
        else:
            description = text(entry.get("description")) or f"{length} {cable_type} cable for {designator}"
        vendor = text(entry.get("vendor")) or records[0].vendor
        notes = f"Manufacturer/vendor: {vendor}. {records[0].source_note}. Cable length from ICD: {length}. Connections: {records[0].endpoints}. {source_note(source_name, records)}"
        if vendor_pn == "TBD":
            notes = f"Manufacturer/vendor: TBD. Vendor part number and cable length are TBD in the ICD. {records[0].source_note}. Connections: {records[0].endpoints}. {source_note(source_name, records)}"
        lines.append(
            {
                "order": (1, min(r.row for r in records)),
                "vendor_part_number": vendor_pn,
                "quantity": len(records),
                "description": description,
                "notes": notes,
                "hyperlink": text(entry.get("canonical_url")),
            }
        )

    lines.sort(key=lambda item: item["order"])
    return lines


def write_workbook(output: Path, lines: list[dict[str, Any]], source_name: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cable BOM"
    wb.properties.title = "Sol Network and Comm Cable Kit BOM"
    wb.properties.subject = f"Cable BOM derived from {source_name}"
    wb.properties.description = "Arena part numbers and drawing-note cells are intentionally blank."

    ws.append(HEADERS)
    for index, line in enumerate(lines, start=1):
        ws.append(
            [
                index,
                None,
                line["vendor_part_number"],
                line["quantity"],
                line["description"],
                None,
            ]
        )

    for cell_obj in ws[1]:
        cell_obj.font = Font(bold=True)
        cell_obj.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 24

    for row in range(2, ws.max_row + 1):
        for column in range(1, 7):
            ws.cell(row, column).alignment = Alignment(vertical="top", wrap_text=column == 5)
        ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row, 4).alignment = Alignment(horizontal="center", vertical="top")
        description_length = len(text(ws.cell(row, 5).value))
        ws.row_dimensions[row].height = min(45, max(20, 15 * (1 + description_length // 60)))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{ws.max_row}"
    widths = {"A": 13, "B": 22, "C": 23, "D": 13, "E": 62, "F": 36}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    ws.print_area = f"A1:F{ws.max_row}"

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def validate_output(output: Path, expected_quantity: int, expected_lines: int) -> dict[str, Any]:
    wb = load_workbook(output, data_only=False)
    if wb.sheetnames != ["Cable BOM"]:
        raise AssertionError(f"Unexpected worksheets: {wb.sheetnames}")
    ws = wb["Cable BOM"]
    actual_headers = [ws.cell(1, column).value for column in range(1, 7)]
    if actual_headers != HEADERS:
        raise AssertionError(f"Unexpected headers: {actual_headers}")
    actual_lines = ws.max_row - 1
    actual_quantity = sum(int(ws.cell(row, 4).value) for row in range(2, ws.max_row + 1))
    if actual_lines != expected_lines:
        raise AssertionError(f"Expected {expected_lines} BOM lines, found {actual_lines}")
    if actual_quantity != expected_quantity:
        raise AssertionError(f"Expected quantity {expected_quantity}, found {actual_quantity}")
    if ws.freeze_panes != "A2" or ws.auto_filter.ref != f"A1:F{ws.max_row}" or ws.tables:
        raise AssertionError("Workbook formatting validation failed")
    if not all(cell.font.bold for cell in ws[1]):
        raise AssertionError("Header row must be bold")
    if any(ws.cell(row, 6).value not in (None, "") for row in range(2, ws.max_row + 1)):
        raise AssertionError("Notes column must be blank")
    if any(cell.hyperlink is not None for row in ws.iter_rows() for cell in row):
        raise AssertionError("Workbook must not contain hyperlinks")
    if any(cell.comment is not None for row in ws.iter_rows() for cell in row):
        raise AssertionError("Workbook must not contain cell comments")
    if any(cell.fill.fill_type is not None for row in ws.iter_rows() for cell in row):
        raise AssertionError("Workbook must not contain color fills")
    return {
        "output": str(output),
        "bom_lines": actual_lines,
        "total_quantity": actual_quantity,
        "tbd_vendor_part_lines": sum(ws.cell(row, 3).value == "TBD" for row in range(2, ws.max_row + 1)),
        "blank_arena_part_numbers": sum(ws.cell(row, 2).value in (None, "") for row in range(2, ws.max_row + 1)),
        "blank_notes": sum(ws.cell(row, 6).value in (None, "") for row in range(2, ws.max_row + 1)),
        "hyperlinks": sum(cell.hyperlink is not None for row in ws.iter_rows() for cell in row),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Readable Network and Comm ICD .xlsx file")
    parser.add_argument("output", type=Path, help="Destination BOM .xlsx file")
    args = parser.parse_args()

    if not args.source.is_file():
        raise SystemExit(f"Source workbook not found: {args.source}")
    wb = load_workbook(args.source, data_only=False, read_only=False)
    for sheet_name in ("Interrack Cabling", "Non-Ethernet Communication"):
        if sheet_name not in wb.sheetnames:
            raise SystemExit(f"Required worksheet not found: {sheet_name}")
    interrack = read_interrack(wb["Interrack Cabling"])
    non_ethernet = read_non_ethernet(wb["Non-Ethernet Communication"])
    lines = build_lines(args.source.name, interrack, non_ethernet, load_catalog())
    write_workbook(args.output, lines, args.source.name)
    summary = validate_output(args.output, len(interrack) + len(non_ethernet), len(lines))
    summary.update({"interrack_source_rows": len(interrack), "non_ethernet_source_rows": len(non_ethernet)})
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
